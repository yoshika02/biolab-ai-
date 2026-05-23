from docx import Document

docx_path = r"C:\Users\yoshi\Downloads\biolab-ai-main\BioLab_AI_Product_Spec.docx"
output_path = r"C:\Users\yoshi\Downloads\biolab-ai-main\biolab-ai-main\PRODUCT_SPEC.md"

try:
    doc = Document(docx_path)
    
    with open(output_path, 'w', encoding='utf-8') as f:
        for para in doc.paragraphs:
            if para.text.strip():
                f.write(para.text + '\n')
        
        # Also extract from tables
        for table in doc.tables:
            f.write('\n| ')
            for cell in table.rows[0].cells:
                f.write(cell.text + ' | ')
            f.write('\n|')
            for cell in table.rows[0].cells:
                f.write(' --- |')
            f.write('\n')
            
            for row in table.rows[1:]:
                f.write('| ')
                for cell in row.cells:
                    f.write(cell.text + ' | ')
                f.write('\n')
    
    print(f"Successfully extracted to {output_path}")
except Exception as e:
    print(f"Error: {e}")
